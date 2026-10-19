#!/usr/bin/env python3
"""Korrekturen am TypeScript Primer (Originalfolien 1-69) von presentation_angular21.pptx.

Vorverarbeitung VOR build_deck.py, auf dem unveränderten Original (309 Folien):

    git show baf0614:presentation_angular21.pptx > /tmp/original.pptx
    python3 -I folien_neu/deck/ts_korrekturen.py /tmp/original.pptx /tmp/original_korrigiert.pptx
    python3 -I folien_neu/deck/build_deck.py /tmp/original_korrigiert.pptx /tmp/neu.pptx

Alle Code-Korrekturen wurden vorher mit dem TypeScript-Compiler (strict) geprüft:
 - Folie 3: Tippfehler im npm-Befehl (concurrently, Bindestrich bei -g)
 - Folie 13: Contextual-Type-Beispiel zeigt jetzt wirklich einen Fehler (.kangaroo)
 - Folie 19: if (flag) statt if (input)
 - Folie 21: let row statt var row
 - Folien 26/27: toFixed(6), Literale statt new Number/new String, const statt var
 - Folie 38: const fahrzeug = v.pop() mit Hinweis auf den Typ Fahrzeug | undefined
 - Folie 45: Methode mit this.name statt globalem name
 - Folie 56: set-Schlüsselwort beim Setter
 - Folie 60: falsche Aussagen zu "Interface erbt von Klasse" korrigiert
 - Farben: Code war in den Farben des dunklen VS-Code-Themes auf hellem Grund gesetzt (schlecht lesbar),
   jetzt Farben des hellen Themes
 - Sprechtexte der Folien 8, 9, 10, 21, 38, 45, 57, 60, 61 korrigiert bzw. neu geschrieben
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.text.text import _Run

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

# Farben des dunklen VS-Code-Themes -> helles Theme (Dark+ -> Light+)
RECOLOR = {
    'D4D4D4': '000000', '9CDCFE': '001080', '569CD6': '0000FF', '4EC9B0': '267F99',
    'DCDCAA': '795E26', 'CE9178': 'A31515', 'B5CEA8': '098658', 'C586C0': 'AF00DB',
    '6A9955': '008000',
}
DARK = dict(default='D4D4D4', var='9CDCFE', kw='569CD6', type='4EC9B0', fn='DCDCAA', str='CE9178',
            num='B5CEA8', ctl='C586C0', cmt='6A9955')


def iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)


def shape_by_name(slide, name):
    for sh in iter_shapes(slide.shapes):
        if sh.name == name:
            return sh
    raise SystemExit('Folie: Shape %r nicht gefunden' % name)


def para_text(para):
    return ''.join(r.text for r in para.runs)


def find_para(slide, startswith=None, contains=None, shape=None):
    hits = []
    for sh in iter_shapes(slide.shapes):
        if not sh.has_text_frame or (shape and sh.name != shape):
            continue
        for para in sh.text_frame.paragraphs:
            t = para_text(para)
            if (startswith and t.strip().startswith(startswith)) or (contains and contains in t):
                hits.append(para)
    if len(hits) != 1:
        raise SystemExit('Absatz nicht eindeutig gefunden (%s/%s): %d Treffer' % (startswith, contains, len(hits)))
    return hits[0]


def set_color(run_el, hex_):
    _Run(run_el, None).font.color.rgb = RGBColor.from_string(hex_)


def rewrite_para(para, tokens):
    """Absatz aus (Text, Farbe)-Tokens neu aufbauen; Schrift/Größe vom ersten Run."""
    runs = para.runs
    proto = copy.deepcopy(runs[0]._r)
    p = para._p
    for ch in list(p):
        if ch.tag in ('{%s}r' % A, '{%s}br' % A, '{%s}fld' % A):
            p.remove(ch)
    end = p.find('{%s}endParaRPr' % A)
    for text, color in tokens:
        r = copy.deepcopy(proto)
        r.find('{%s}t' % A).text = text
        rp = r.find('{%s}rPr' % A)
        for k in ('err', 'dirty'):
            rp.attrib.pop(k, None)
        set_color(r, color)
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)


def replace_in_runs(slide, old, new, exact=False):
    n = 0
    for sh in iter_shapes(slide.shapes):
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if (exact and r.text == old) or (not exact and old in r.text):
                    r.text = r.text.replace(old, new) if not exact else new
                    n += 1
    if n == 0:
        raise SystemExit('Text %r nicht gefunden' % old)
    return n


def replace_var_keyword(slide):
    return replace_in_runs(slide, 'var', 'const', exact=True)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def main(src, dst):
    prs = Presentation(src)
    assert len(prs.slides) == 309, 'Erwartet wird das Original mit 309 Folien'
    S = {i: s for i, s in enumerate(prs.slides, 1)}

    # Folie 3: Tippfehler im Befehl
    replace_in_runs(S[3], 'concurrenty', 'concurrently')
    replace_in_runs(S[3], ' –', ' -', exact=True)

    # Folie 13: Contextual Type zeigt jetzt tatsaechlich einen Fehler
    replace_in_runs(S[13], '.button);  //<- Error, ', '.kangaroo);  // <- Fehler: gibt es nicht')

    # Folie 19: Parameter heisst flag
    replace_in_runs(S[19], 'input', 'flag', exact=True)

    # Folie 21: let statt var
    replace_in_runs(S[21], '\tvar', '\tlet', exact=True)

    # Folien 26/27: Beispiele ohne Wrapper-Objekte, konsistent const
    replace_var_keyword(S[26])
    replace_var_keyword(S[27])
    rewrite_para(find_para(S[26], contains='numToFixed = new Number'),
                 [('const', DARK['kw']), (' ', DARK['default']), ('numToFixed', DARK['var']),
                  (' = ', DARK['default']), ('123.456', DARK['num']), (';', DARK['default'])])
    rewrite_para(find_para(S[26], contains='numTo = new Number'),
                 [('const', DARK['kw']), (' ', DARK['default']), ('numTo', DARK['var']),
                  (' = ', DARK['default']), ('123.456', DARK['num']), (';', DARK['default'])])
    replace_in_runs(S[26], '4', '6', exact=True)  # toFixed(4) -> toFixed(6), passend zur Beschriftung
    rewrite_para(find_para(S[27], contains='new String'),
                 [('const', DARK['kw']), (' ', DARK['default']), ('str1', DARK['var']),
                  (' = ', DARK['default']), ('"The TS String"', DARK['str']), (';', DARK['default'])])

    # Folie 38: pop() liefert im strikten Modus Fahrzeug | undefined
    replace_var_keyword(S[38])
    rewrite_para(find_para(S[38], contains='fahrzeug:Fahrzeug'),
                 [('const', DARK['kw']), (' ', DARK['default']), ('fahrzeug', DARK['var']),
                  (' = ', DARK['default']), ('v', DARK['var']), ('.', DARK['default']),
                  ('pop', DARK['fn']), ('();', DARK['default']),
                  ('   // Typ: Fahrzeug | undefined', DARK['cmt'])])
    replace_in_runs(S[38], 'Fahrzeug ct =', 'const ct =')

    # Folie 45: Methode mit this.name
    replace_var_keyword(S[45])
    replace_in_runs(S[45], 'mit var', 'mit Variable')   # Titel: "mit Variable" statt "mit var"
    rewrite_para(find_para(S[45], startswith='suhlen'),
                 [('suhlen', DARK['fn']), ('() { ', DARK['default']), ('return', DARK['ctl']),
                  (' ', DARK['default']), ('this', DARK['kw']), ('.', DARK['default']),
                  ('name', DARK['var']), (' + ', DARK['default']),
                  ('" ist gerne im Dreck"', DARK['str']), ('; }', DARK['default'])])

    # Folie 56: Setter braucht das Schluesselwort set
    p56 = find_para(S[56], startswith='name(new_name')
    first = copy.deepcopy(p56.runs[0]._r)
    first.find('{%s}t' % A).text = 'set '
    set_color(first, DARK['kw'])
    p56.runs[0]._r.addprevious(first)

    # Folie 60: falsche Aussagen zu Interface/Klasse ersetzen
    set_para = lambda para, text: [setattr(para.runs[0], 'text', text)] + [r._r.getparent().remove(r._r) for r in para.runs[1:]]
    set_para(find_para(S[60], startswith='Ein Interface kann von einer Klasse erben'),
             'Ein Interface kann eine Klasse erweitern (kommt selten vor)')
    set_para(find_para(S[60], startswith='Allerdings darf die Klasse'),
             'Es übernimmt die Mitglieder der Klasse, aber nicht deren Implementierung')

    # Farben des hellen Themes statt dunklem (Folien 1-69)
    n = 0
    for i in range(1, 70):
        for sh in iter_shapes(S[i].shapes):
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    c = r.font.color
                    if c is not None and c.type == 1 and str(c.rgb) in RECOLOR:
                        c.rgb = RGBColor.from_string(RECOLOR[str(c.rgb)])
                        n += 1
    print('Farben umgestellt:', n, 'Runs')

    # Sprechtexte
    set_notes(S[8], (
        'Folie: „TypeScript – Einordnung“.\n\n'
        'Das Bild zeigt eine Rechnung aus Kreisen. Der erste Kreis steht für „ECMAScript 5“, eine ältere Version von JavaScript. '
        'Der zweite Kreis steht für „ECMAScript 6“, die neuere Version mit zusätzlichen Möglichkeiten wie Klassen und „let“. '
        'Der dritte Kreis heißt „Additional Features“, auf Deutsch zusätzliche Eigenschaften. Das Ergebnis der Rechnung ist TypeScript.\n\n'
        '„ECMAScript“ ist der offizielle Name des Standards, nach dem JavaScript funktioniert. Die Aussage der Folie: TypeScript '
        'enthält alles von JavaScript und fügt vor allem Typen hinzu. Der Browser versteht TypeScript nicht direkt. Deshalb wird der '
        'Code vorher in normales JavaScript übersetzt.'))
    set_notes(S[9], (
        'Folie: „Warum TypeScript“.\n\n'
        'Die Folie nennt vier Gründe in vier Feldern. Erstens „OO Concepts“: TypeScript unterstützt die Ideen der objektorientierten '
        'Programmierung, also Klassen, Vererbung und Schnittstellen. Zweitens „Supports Type Definitions“: Man kann für Variablen, '
        'Funktionen und Objekte festlegen, welche Art von Werten erlaubt ist. Drittens „Strong static Type“: Die Typen werden streng '
        'geprüft, und zwar schon beim Schreiben und Übersetzen des Codes, nicht erst, wenn das Programm läuft. Viertens „Compilation“: '
        'Der Code wird vor der Ausführung vom Compiler geprüft und in JavaScript übersetzt.\n\n'
        'Im Alltag bedeutet das: Der Editor findet Fehler früh und macht Vorschläge beim Schreiben, und größere Projekte bleiben besser '
        'wartbar. Das hilft bei der Kunden-App im Kurs, denn auch Angular selbst ist in TypeScript geschrieben.'))
    set_notes(S[10], (
        'Folie: „Type Hierarchy“.\n\n'
        'Die Folie zeigt drei Ebenen untereinander. Ganz oben steht „any“. Das ist ein Typ, der jeden Wert erlaubt. Darunter stehen die '
        'eingebauten Typen, englisch „build-in“, zum Beispiel Zahl, Text und Wahrheitswert. Ganz unten stehen die selbst definierten '
        'Typen, englisch „user-defined“, zum Beispiel eigene Klassen und Schnittstellen.\n\n'
        'Das Bild ist stark vereinfacht. Wichtig ist die Idee: Jeder konkrete Typ ist ein Sonderfall von „any“. Deshalb darf man jeden '
        'Wert einer „any“-Variable zuweisen. Genau das ist aber gefährlich, denn dann prüft der Compiler nichts mehr. In eigenem Code '
        'sollte man „any“ deshalb möglichst vermeiden.'))
    set_notes(S[21], (
        'Folie: „Block Scoping 3“ – die Lösung des Problems von „Var-Scoping Macken 2“ (verschachtelte Schleifen) mit „let“ statt „var“.\n\n'
        'Bei „Var-Scoping Macken 2“ benutzten beide Schleifen mit „var“ dieselbe Variable „i“. Die innere Schleife hat dadurch den Zähler '
        'der äußeren Schleife überschrieben. Mit „let“ hat jede Schleife ihre eigene Variable „i“, die nur innerhalb dieser Schleife gilt. '
        'Die Summe wird deshalb richtig berechnet.\n\n'
        'Auch das „setTimeout“-Beispiel von der Folie davor wäre mit „let“ richtig: Jeder Durchlauf bekommt seine eigene Kopie von „i“, '
        'und es erscheinen die Zahlen 0 bis 9. Das ist einer der Hauptgründe, warum „let“ heute Standard ist.'))
    set_notes(S[38], S[38].notes_slide.notes_text_frame.text +
              '\n\nIm strikten Modus gibt „pop“ den Typ „Fahrzeug oder undefined“ zurück, weil die Liste leer sein könnte. Das zeigt der '
              'Kommentar im Code. Man müsste das Ergebnis also vor der Benutzung auf „undefined“ prüfen.')
    set_notes(S[45], S[45].notes_slide.notes_text_frame.text +
              '\n\nDie Methode „suhlen“ greift mit „this.name“ auf die Eigenschaft „name“ desselben Objekts zu. Das Wort „this“ bedeutet: '
              'dieses Objekt.')
    set_notes(S[57], (
        'Folie: „Interfaces und Klassen im Diagramm“.\n\n'
        'Das Bild ist ein Klassendiagramm in der Schreibweise „UML“. Oben stehen zwei Schnittstellen mit den Namen „Flushable“ und '
        '„Writeable“. Darunter steht eine Klasse „Writer“. Die gestrichelten Pfeile von „Writer“ zu den beiden Schnittstellen bedeuten: '
        '„Writer“ setzt beide Schnittstellen um. Die Klasse muss deshalb alle dort verlangten Methoden anbieten, hier zum Beispiel '
        '„write“ und „flush“.\n\n'
        'Das Beispiel stammt ursprünglich aus der Java-Welt. Die Idee gilt aber genauso in TypeScript: Eine Klasse kann mehrere '
        'Schnittstellen gleichzeitig erfüllen. Auf der nächsten Folie sehen wir das als TypeScript-Code.'))
    set_notes(S[60], (
        'Folie: „Klassen: Vererbung“.\n\n'
        'Eine Klasse kann von einer anderen Klasse erben. Dafür gibt es das Wort „extends“. Anders als bei „implements“ bekommt die '
        'Klasse dabei die echte Umsetzung der Oberklasse und muss nicht alles selbst schreiben. Methoden können „abstract“ sein: Die '
        'Basisklasse gibt nur die Signatur vor, jede konkrete Unterklasse muss die Methode selbst umsetzen. Eine Klasse mit abstrakten '
        'Methoden kann man nicht mit „new“ erzeugen.\n\n'
        'Eine Klasse kann außerdem eine oder mehrere Schnittstellen mit „implements“ umsetzen. Ein Interface kann sogar eine Klasse '
        'erweitern. Es übernimmt dann deren Mitglieder, aber nicht deren Umsetzung. Das kommt selten vor.'))
    set_notes(S[61], (
        'Codefolie: „Vererbung im Beispiel“.\n\n'
        'Die Klasse „Schwein“ setzt das Interface „Schmutzfink“ um und hat die Methode „suhlen“. Die Klasse „Mastschwein“ erbt mit '
        '„extends“ von „Schwein“. Sie bekommt dadurch die Methode „suhlen“ automatisch und hat zusätzlich ein eigenes Attribut '
        '„genmodifikation“.\n\n'
        'Im Konstruktor von „Mastschwein“ ruft „super(name)“ den Konstruktor der Oberklasse auf, damit der Name gesetzt wird. „super“ '
        'muss im Konstruktor einer Unterklasse zuerst aufgerufen werden, bevor man „this“ benutzt.'))

    prs.save(dst)
    print('geschrieben:', dst)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
