#!/usr/bin/env python3
"""Fuegt neue Folien in presentation_angular21.pptx ein (siehe SCHEMA.md / README.md).

Aufruf:
    python3 build_deck.py <eingabe.pptx> <ausgabe.pptx> [--content-dir DIR] [--quiet]

Die Eingabedatei wird nie veraendert; die Ausgabe wird immer frisch aus der
unveraenderten Eingabe erzeugt (idempotent).
"""
import argparse
import copy
import glob
import json
import math
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# Vorlagenfolien (1-basiert, Originaldatei)
TEMPLATE_SLIDE = {"content": 256, "code": 160, "exercise": 258, "divider": 269}
TYPES = tuple(TEMPLATE_SLIDE)

MAX_LINES = {"content": 9, "exercise": 9, "code": 18}
MAX_CHARS = {"content": 95, "exercise": 95, "code": 90}

_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


class BuildError(Exception):
    pass


# --------------------------------------------------------------------------
# Hilfsfunktionen
# --------------------------------------------------------------------------
def clean(text):
    """Entfernt unzulaessige XML-Zeichen, normalisiert Zeilenumbrueche."""
    if text is None:
        return ""
    text = str(text).replace("\r\n", "\n").replace("\r", "\n")
    return _INVALID_XML.sub("", text)


def clean_notes(text):
    text = clean(text).replace("\t", " ").strip("\n")
    return re.sub(r"\n{3,}", "\n\n", text)


def sub(parent, tag, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def wrap_lines(text, width_pt, size_pt, char_factor):
    """Schaetzt die Zahl der Zeilen beim Umbruch (Wortumbruch)."""
    cpl = max(8, int(width_pt / (size_pt * char_factor)))
    if not text:
        return 1
    lines, cur = 1, 0
    for word in text.split(" "):
        wl = len(word)
        if cur == 0:
            cur = wl
        elif cur + 1 + wl <= cpl:
            cur += 1 + wl
        else:
            lines += 1
            cur = wl
        while cur > cpl:  # sehr langes Wort
            lines += 1
            cur -= cpl
    return lines


def fit_title(text, sizes, width_pt, height_pt, char_factor, line_factor):
    """Waehlt die groesste Schriftgroesse (in 1/100 pt), bei der der Titel passt."""
    for sz in sizes:
        pt = sz / 100.0
        n = wrap_lines(text, width_pt, pt, char_factor)
        if n * pt * line_factor <= height_pt:
            return sz, True
    return sizes[-1], False


def strip_creation_ids(el):
    """Entfernt a16:creationId-Erweiterungen (sonst doppelte IDs in Kopien)."""
    for ext_lst in list(el.iter(qn("a:extLst"))):
        parent = ext_lst.getparent()
        if parent is None or parent.tag != qn("p:cNvPr"):
            continue
        for ext in list(ext_lst):
            if ext.get("uri") == "{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}":
                ext_lst.remove(ext)
        if len(ext_lst) == 0:
            parent.remove(ext_lst)


def get_rpr(para):
    """Liefert (oder erzeugt) Eigenschaften des ersten Runs als Prototyp."""
    r = para.find(qn("a:r"))
    if r is not None:
        rpr = r.find(qn("a:rPr"))
        if rpr is not None:
            return copy.deepcopy(rpr)
    rpr = etree.Element(qn("a:rPr"))
    rpr.set("lang", "de-DE")
    ppr = para.find(qn("a:pPr"))
    d = ppr.find(qn("a:defRPr")) if ppr is not None else None
    if d is not None:
        for k, v in d.attrib.items():
            rpr.set(k, v)
        for ch in d:
            rpr.append(copy.deepcopy(ch))
    return rpr


def set_sz(rpr, sz, bold=None):
    rpr.set("sz", str(int(sz)))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")


def make_para(proto, text, sz=None, bold=None, lvl=None, spc_bef=None):
    """Neuer Absatz auf Basis eines Prototyp-Absatzes (Formatierung kopiert)."""
    p = etree.Element(qn("a:p"))
    ppr_src = proto.find(qn("a:pPr"))
    ppr = copy.deepcopy(ppr_src) if ppr_src is not None else etree.Element(qn("a:pPr"))
    p.append(ppr)
    if lvl:
        ppr.set("lvl", str(lvl))
    elif "lvl" in ppr.attrib and lvl == 0:
        del ppr.attrib["lvl"]
    d = ppr.find(qn("a:defRPr"))
    if d is not None and sz:
        d.set("sz", str(int(sz)))
        if bold is not None:
            d.set("b", "1" if bold else "0")
    if spc_bef is not None:
        sb = ppr.find(qn("a:spcBef"))
        if sb is not None:
            sb.find(qn("a:spcPts")).set("val", str(int(spc_bef)))
    rpr = get_rpr(proto)
    if sz:
        set_sz(rpr, sz, bold)
    r = sub(p, "a:r")
    r.append(rpr)
    t = sub(r, "a:t")
    t.text = clean(text).replace("\t", "    ")
    end = proto.find(qn("a:endParaRPr"))
    if end is not None:
        e = copy.deepcopy(end)
        if sz:
            e.set("sz", str(int(sz)))
        p.append(e)
    return p


def replace_paragraphs(txbody, paras):
    for p in txbody.findall(qn("a:p")):
        txbody.remove(p)
    for p in paras:
        txbody.append(p)


def ph_type(sp):
    ph = sp.find(".//" + qn("p:ph"))
    if ph is None:
        return None
    return ph.get("type", "body")


def find_title_sp(spTree):
    for sp in spTree.iter(qn("p:sp")):
        if ph_type(sp) in ("title", "ctrTitle"):
            return sp
    return None


def find_body_sp(spTree):
    cands = []
    for sp in spTree.iter(qn("p:sp")):
        t = ph_type(sp)
        if t in ("title", "ctrTitle", "ftr", "sldNum", "dt", "sldImg"):
            continue
        if sp.find(qn("p:txBody")) is None:
            continue
        cands.append((t is not None, sum(len("".join(p.itertext())) for p in sp.iter(qn("a:p"))), sp))
    if not cands:
        return None
    # Platzhalter zuerst, dann der mit dem meisten Text
    cands.sort(key=lambda c: (c[0], c[1]), reverse=True)
    return cands[0][2]


def normalize_bullets(raw, where):
    out = []
    if not isinstance(raw, list):
        raise BuildError("%s: 'bullets' muss eine Liste sein" % where)
    for i, b in enumerate(raw):
        if isinstance(b, str):
            out.append((clean(b), 0))
        elif isinstance(b, (list, tuple)) and len(b) in (1, 2) and isinstance(b[0], str):
            lvl = b[1] if len(b) == 2 else 0
            if lvl not in (0, 1):
                raise BuildError("%s: bullets[%d]: Ebene muss 0 oder 1 sein" % (where, i))
            out.append((clean(b[0]), lvl))
        else:
            raise BuildError("%s: bullets[%d] muss String oder [Text, Ebene] sein" % (where, i))
    return out


# --------------------------------------------------------------------------
# Inhalt laden und validieren
# --------------------------------------------------------------------------
def load_contents(content_dir, n_slides):
    files = sorted(glob.glob(os.path.join(content_dir, "content_*.json")))
    blocks, ops = [], []
    for fpath in files:
        fname = os.path.basename(fpath)
        try:
            with open(fpath, encoding="utf-8") as fh:
                data = json.load(fh)
        except (OSError, ValueError) as exc:
            raise BuildError("%s: JSON nicht lesbar: %s" % (fname, exc))
        if not isinstance(data, dict):
            raise BuildError("%s: Wurzel muss ein Objekt sein" % fname)
        for key in data:
            if key not in ("blocks", "ops"):
                print("WARNUNG: %s: unbekannter Schluessel '%s' ignoriert" % (fname, key))
        for bi, blk in enumerate(data.get("blocks", [])):
            where = "%s: blocks[%d]" % (fname, bi)
            if not isinstance(blk, dict) or "after" not in blk or "slides" not in blk:
                raise BuildError("%s: 'after' und 'slides' erforderlich" % where)
            after = blk["after"]
            if not isinstance(after, int) or isinstance(after, bool) or not 0 <= after <= n_slides:
                raise BuildError("%s: 'after'=%r ist keine gueltige Folie (0..%d)" % (where, after, n_slides))
            order = blk.get("order", 0)
            if not isinstance(order, (int, float)):
                raise BuildError("%s: 'order' muss eine Zahl sein" % where)
            if not isinstance(blk["slides"], list):
                raise BuildError("%s: 'slides' muss eine Liste sein" % where)
            slides = []
            for si, s in enumerate(blk["slides"]):
                swhere = "%s.slides[%d]" % (where, si)
                slides.append(validate_slide(s, swhere, fname))
            blocks.append({"after": after, "order": order, "file": fname, "idx": bi, "slides": slides})
        for oi, op in enumerate(data.get("ops", [])):
            ops.append(validate_op(op, "%s: ops[%d]" % (fname, oi), fname, n_slides))
    return files, blocks, ops


def validate_slide(s, where, fname):
    if not isinstance(s, dict):
        raise BuildError("%s: Folie muss ein Objekt sein" % where)
    typ = s.get("type")
    if typ not in TYPES:
        raise BuildError("%s: unbekannter type %r (erlaubt: %s)" % (where, typ, ", ".join(TYPES)))
    title = clean(s.get("title", "")).strip()
    if not title:
        raise BuildError("%s: 'title' fehlt" % where)
    out = {"type": typ, "title": title, "notes": clean_notes(s.get("notes", "")), "file": fname, "where": where}
    if typ in ("content", "exercise"):
        if "bullets" not in s:
            raise BuildError("%s: 'bullets' fehlt (type %s)" % (where, typ))
        out["bullets"] = normalize_bullets(s["bullets"], where)
    elif typ == "code":
        if not isinstance(s.get("code"), str):
            raise BuildError("%s: 'code' (String) fehlt (type code)" % where)
        out["code"] = clean(s["code"]).replace("\t", "    ").rstrip("\n")
    return out


def _num(op, key, where, n_slides, allow_zero=False):
    v = op.get(key)
    lo = 0 if allow_zero else 1
    if not isinstance(v, int) or isinstance(v, bool) or not lo <= v <= n_slides:
        raise BuildError("%s: '%s'=%r ist keine gueltige Folie (%d..%d)" % (where, key, v, lo, n_slides))
    return v


def validate_op(op, where, fname, n_slides):
    if not isinstance(op, dict):
        raise BuildError("%s: Operation muss ein Objekt sein" % where)
    kind = op.get("op")
    o = {"op": kind, "file": fname, "where": where}
    if kind == "move":
        o["from"] = _num(op, "from", where, n_slides)
        o["to"] = _num(op, "to", where, n_slides)
        o["after"] = _num(op, "after", where, n_slides, allow_zero=True)
        if o["from"] > o["to"]:
            raise BuildError("%s: 'from' (%d) groesser als 'to' (%d)" % (where, o["from"], o["to"]))
        if o["from"] <= o["after"] <= o["to"]:
            raise BuildError("%s: Ziel 'after'=%d liegt im verschobenen Bereich %d..%d"
                             % (where, o["after"], o["from"], o["to"]))
    elif kind == "delete":
        sl = op.get("slides")
        if not isinstance(sl, list) or not sl:
            raise BuildError("%s: 'slides' muss eine nichtleere Liste sein" % where)
        o["slides"] = []
        for v in sl:
            if not isinstance(v, int) or isinstance(v, bool) or not 1 <= v <= n_slides:
                raise BuildError("%s: unbekannte Folie %r in 'slides' (1..%d)" % (where, v, n_slides))
            o["slides"].append(v)
    elif kind == "edit":
        o["slide"] = _num(op, "slide", where, n_slides)
        if not any(k in op for k in ("title", "bullets", "notes", "code")):
            raise BuildError("%s: edit ohne title/bullets/notes/code" % where)
        if "title" in op:
            o["title"] = clean(op["title"]).strip()
            if not o["title"]:
                raise BuildError("%s: leerer 'title'" % where)
        if "bullets" in op:
            o["bullets"] = normalize_bullets(op["bullets"], where)
        if "code" in op:
            if not isinstance(op["code"], str):
                raise BuildError("%s: 'code' muss String sein" % where)
            o["code"] = clean(op["code"]).replace("\t", "    ").rstrip("\n")
        if "notes" in op:
            o["notes"] = clean_notes(op["notes"])
    elif kind == "replace_text":
        o["slide"] = _num(op, "slide", where, n_slides)
        if not isinstance(op.get("old"), str) or not op["old"] or not isinstance(op.get("new"), str):
            raise BuildError("%s: 'old' (nichtleer) und 'new' (String) erforderlich" % where)
        o["old"], o["new"] = op["old"], op["new"]
        o["include_notes"] = bool(op.get("include_notes", False))
    else:
        raise BuildError("%s: unbekannte Operation %r (move/delete/edit/replace_text)" % (where, kind))
    return o


# --------------------------------------------------------------------------
# Neue Folien erzeugen
# --------------------------------------------------------------------------
class Templates:
    """Haelt Kopien der Vorlagen-Shapes (vor allen Aenderungen gezogen)."""

    def __init__(self, prs):
        slides = list(prs.slides)
        self.data = {}
        for typ, num in TEMPLATE_SLIDE.items():
            if num > len(slides):
                raise BuildError("Eingabedatei hat keine Vorlagenfolie %d (type %s)" % (num, typ))
            sl = slides[num - 1]
            tree = copy.deepcopy(sl._element.cSld.spTree)
            shapes = [copy.deepcopy(ch) for ch in tree
                      if ch.tag in (qn("p:sp"), qn("p:pic"), qn("p:cxnSp"), qn("p:grpSp"), qn("p:graphicFrame"))]
            for s in shapes:
                strip_creation_ids(s)
            timing = sl._element.find(qn("p:timing"))
            rels = {}
            for rid, rel in sl.part.rels.items():
                if rel.reltype.endswith(("/slideLayout", "/notesSlide")) or rel.is_external:
                    continue
                rels[rid] = rel
            self.data[typ] = {"layout": sl.slide_layout, "shapes": shapes, "slide": sl,
                              "timing": copy.deepcopy(timing) if timing is not None else None,
                              "rels": rels}


def new_slide_from_template(prs, tpl, typ):
    t = tpl.data[typ]
    slide = prs.slides.add_slide(t["layout"])
    tree = slide.shapes._spTree
    for ch in list(tree):
        if ch.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.remove(ch)
    rid_map = {}
    for rid, rel in t["rels"].items():
        rid_map[rid] = slide.part.relate_to(rel.target_part, rel.reltype)
    for sp in t["shapes"]:
        new = copy.deepcopy(sp)
        for el in new.iter():
            for k in list(el.attrib):
                if k.startswith("{%s}" % R) and el.get(k) in rid_map:
                    el.set(k, rid_map[el.get(k)])
        tree.append(new)
    if t["timing"] is not None:
        old = slide._element.find(qn("p:timing"))
        if old is not None:
            slide._element.remove(old)
        slide._element.append(copy.deepcopy(t["timing"]))
    return slide, tree


def build_title(tree, title, sizes, width_pt, height_pt, char_factor, line_factor, warns, ctx):
    sp = find_title_sp(tree)
    body = sp.find(qn("p:txBody"))
    proto = body.find(qn("a:p"))
    sz, ok = fit_title(title, sizes, width_pt, height_pt, char_factor, line_factor)
    if not ok:
        warns.append("%s: Titel sehr lang (%d Zeichen), kleinste Schrift %d pt genutzt - bitte kuerzen"
                     % (ctx, len(title), sz // 100))
    elif sz != sizes[0]:
        warns.append("%s: Titel lang (%d Zeichen), Schrift auf %d pt (Vorlage %d pt) verkleinert"
                     % (ctx, len(title), sz // 100, sizes[0] // 100))
    replace_paragraphs(body, [make_para(proto, title, sz=sz)])


def build_lines_body(tree, bullets, kind, warns, ctx):
    """Aufzaehlung (content/exercise) im Stil der Folien 256/258."""
    sp = find_body_sp(tree)
    body = sp.find(qn("p:txBody"))
    protos = body.findall(qn("a:p"))
    proto = protos[0]
    bpr = body.find(qn("a:bodyPr"))
    if bpr is not None:
        bpr.set("anchor", "t")  # oben beginnen, auch bei kurzen Listen
    W, H = 810.0, 330.0
    CF, PITCH = 0.50, 1.35
    chosen = 1.0
    for scale in (1.0, 0.9, 0.8, 0.7, 0.6):
        total = 0.0
        for text, lvl in bullets:
            sz = (2000 if lvl == 0 else 1700) * scale / 100.0
            total += wrap_lines(text, W - (30 if lvl else 0), sz, CF) * sz * PITCH
        chosen = scale
        if total <= H:
            break
    if chosen < 1.0:
        warns.append("%s: Inhalt lang, Schrift auf %d%% verkleinert (Vorlage 20/17 pt)" % (ctx, round(chosen * 100)))
    paras = []
    for i, (text, lvl) in enumerate(bullets):
        bold = False
        if kind == "exercise" and lvl == 0:
            bold = (i == 0) or bool(re.match(r"^\s*\d+[.)]", text)) or text.startswith("Musterlösung")
        sz = (2000 if lvl == 0 else 1700) * chosen
        paras.append(make_para(proto, text, sz=sz, bold=bold, lvl=lvl))
    replace_paragraphs(body, paras)


def build_code_body(tree, code, warns, ctx):
    sp = find_body_sp(tree)
    body = sp.find(qn("p:txBody"))
    proto = body.find(qn("a:p"))
    lines = code.split("\n")
    W, H = 810.0, 335.0
    maxlen = max((len(l) for l in lines), default=1)
    steps = [(2000, 730), (1800, 300), (1600, 200), (1400, 100), (1200, 0), (1000, 0)]
    chosen = steps[-1]
    fits = False
    for sz, spc in steps:
        pitch = sz / 100.0 * 1.08 + spc / 100.0
        if len(lines) * pitch <= H and maxlen * 0.6 * sz / 100.0 <= W:
            chosen, fits = (sz, spc), True
            break
    if chosen[0] != 2000:
        warns.append("%s: Code lang (%d Zeilen, max. %d Zeichen), Schrift auf %d pt (Vorlage 20 pt) verkleinert"
                     % (ctx, len(lines), maxlen, chosen[0] // 100))
    if not fits:
        warns.append("%s: Code passt auch bei %d pt nicht sicher auf die Folie - bitte kuerzen/teilen"
                     % (ctx, chosen[0] // 100))
    replace_paragraphs(body, [make_para(proto, l, sz=chosen[0], spc_bef=chosen[1]) for l in lines])


def write_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        raise BuildError("Notizseite ohne Textplatzhalter")
    tf.text = text


def build_slide(prs, tpl, spec, warns):
    typ = spec["type"]
    ctx = "%s (%s '%s')" % (spec["where"], typ, spec["title"][:40])
    slide, tree = new_slide_from_template(prs, tpl, typ)
    if typ in ("content", "exercise"):
        build_title(tree, spec["title"], [3600, 3200, 2800, 2400], 800, 100, 0.56, 1.2, warns, ctx)
        build_lines_body(tree, spec["bullets"], typ, warns, ctx)
    elif typ == "code":
        build_title(tree, spec["title"], [4400, 3600, 3200, 2800, 2400], 800, 100, 0.48, 1.08, warns, ctx)
        build_code_body(tree, spec["code"], warns, ctx)
    else:  # divider
        build_title(tree, spec["title"], [7200, 6000, 5400, 4800, 4000, 3600], 810, 185, 0.46, 1.08, warns, ctx)
    if spec["notes"]:
        write_notes(slide, spec["notes"])
    else:
        warns.append("%s: keine Notizen (Sprechtext fehlt)" % ctx)
        write_notes(slide, "")
    # Groessenwarnungen
    if typ in MAX_LINES:
        items = [b[0] for b in spec["bullets"]] if typ != "code" else spec["code"].split("\n")
        if len(items) > MAX_LINES[typ]:
            warns.append("%s: %d Zeilen (Grenze %d)" % (ctx, len(items), MAX_LINES[typ]))
        long_ = [i for i in items if len(i) > MAX_CHARS[typ]]
        if long_:
            warns.append("%s: %d Zeile(n) laenger als %d Zeichen (laengste %d)"
                         % (ctx, len(long_), MAX_CHARS[typ], max(len(i) for i in long_)))
    if len(spec["notes"]) > 3500:
        warns.append("%s: sehr lange Notizen (%d Zeichen)" % (ctx, len(spec["notes"])))
    return slide


# --------------------------------------------------------------------------
# edit / replace_text auf Originalfolien
# --------------------------------------------------------------------------
def apply_edit(slide, op, warns):
    tree = slide.shapes._spTree
    ctx = op["where"]
    if "title" in op:
        sp = find_title_sp(tree)
        if sp is None:
            raise BuildError("%s: Folie hat keinen Titelplatzhalter" % ctx)
        body = sp.find(qn("p:txBody"))
        ps = body.findall(qn("a:p"))
        proto = ps[0] if ps else etree.Element(qn("a:p"))
        replace_paragraphs(body, [make_para(proto, op["title"])])
    lines = None
    if "bullets" in op:
        lines = op["bullets"]
    elif "code" in op:
        lines = [(l, 0) for l in op["code"].split("\n")]
    if lines is not None:
        sp = find_body_sp(tree)
        if sp is None:
            raise BuildError("%s: Folie hat keinen Textkoerper fuer bullets/code" % ctx)
        body = sp.find(qn("p:txBody"))
        ps = body.findall(qn("a:p"))
        texty = [p for p in ps if "".join(p.itertext()).strip()] or ps
        if not texty:
            texty = [etree.Element(qn("a:p"))]
        proto0 = texty[0]
        proto1 = texty[-1] if len(texty) > 1 else texty[0]
        paras = []
        for text, lvl in lines:
            if lvl == 0:
                paras.append(make_para(proto0, text))
            else:
                paras.append(make_para(proto1, text, lvl=(1 if proto1 is proto0 else None)))
        replace_paragraphs(body, paras)
        if len(lines) > 9 and "code" not in op:
            warns.append("%s: %d Zeilen (Grenze 9)" % (ctx, len(lines)))
        if any(len(t) > 95 for t, _ in lines) and "code" not in op:
            warns.append("%s: Zeile(n) laenger als 95 Zeichen" % ctx)
    if "notes" in op:
        if not op["notes"]:
            warns.append("%s: edit mit leeren Notizen" % ctx)
        write_notes(slide, op["notes"])


def _replace_in_paragraph(p, old, new):
    """Ersetzt Text in einem a:p; run-uebergreifend, Formatierung des ersten Runs bleibt."""
    runs = p.findall(qn("a:r"))
    count = 0
    for r in runs:
        t = r.find(qn("a:t"))
        if t is not None and t.text and old in t.text:
            count += t.text.count(old)
            t.text = t.text.replace(old, new)
    full = "".join((r.find(qn("a:t")).text or "") for r in runs if r.find(qn("a:t")) is not None)
    if old in full and len(runs) > 1:  # Treffer ueber Run-Grenzen
        n = full.count(old)
        newfull = full.replace(old, new)
        first = True
        for r in runs:
            t = r.find(qn("a:t"))
            if t is None:
                continue
            if first:
                t.text, first = newfull, False
            else:
                t.text = ""
        count += n
    return count


def _iter_paragraphs(root):
    return root.iter(qn("a:p"))


def apply_replace(slide, op):
    total = 0
    for p in _iter_paragraphs(slide._element):
        total += _replace_in_paragraph(p, op["old"], op["new"])
    if op["include_notes"] and slide.has_notes_slide:
        for p in _iter_paragraphs(slide.notes_slide._element):
            total += _replace_in_paragraph(p, op["old"], op["new"])
    return total


# --------------------------------------------------------------------------
# Hauptablauf
# --------------------------------------------------------------------------
def build(in_path, out_path, content_dir, quiet=False):
    if not os.path.isfile(in_path):
        raise BuildError("Eingabedatei nicht gefunden: %s" % in_path)
    if os.path.abspath(in_path) == os.path.abspath(out_path):
        raise BuildError("Ausgabe darf nicht gleich Eingabe sein (kein In-place): %s" % out_path)
    prs = Presentation(in_path)
    n = len(prs.slides)
    warns = []
    files, blocks, ops = load_contents(content_dir, n)
    tpl = Templates(prs)

    sld_id_lst = prs.slides._sldIdLst
    orig_ids = list(sld_id_lst)          # sldId-Elemente in Originalreihenfolge
    orig_slides = list(prs.slides)
    order = list(range(1, n + 1))        # Eintraege: int (Original) oder ("new", slide)
    deleted = set()
    stats = {"inserted": 0, "moved": 0, "deleted": 0, "edited": 0, "replaced": 0}
    edited, moved_set = set(), set()

    def alive(num, ctx):
        if num in deleted:
            raise BuildError("%s: Folie %d wurde geloescht" % (ctx, num))

    # 1. delete
    for op in (o for o in ops if o["op"] == "delete"):
        for num in op["slides"]:
            deleted.add(num)
    # 2. move
    for op in (o for o in ops if o["op"] == "move"):
        items = [x for x in order if isinstance(x, int) and op["from"] <= x <= op["to"]]
        order = [x for x in order if x not in items]
        pos = 0 if op["after"] == 0 else order.index(op["after"]) + 1
        order[pos:pos] = items
        for x in items:
            if x not in deleted:
                moved_set.add(x)
    stats["moved"] = len(moved_set)
    # 3. edit / replace_text
    for op in ops:
        if op["op"] not in ("edit", "replace_text"):
            continue
        alive(op["slide"], op["where"] + " (%s)" % op["file"])
        sl = orig_slides[op["slide"] - 1]
        if op["op"] == "edit":
            try:
                apply_edit(sl, op, warns)
            except BuildError as exc:
                raise BuildError("%s: %s" % (op["file"], exc))
            edited.add(op["slide"])
        else:
            cnt = apply_replace(sl, op)
            if cnt == 0:
                warns.append("%s: replace_text auf Folie %d: '%s' nicht gefunden"
                             % (op["where"], op["slide"], op["old"]))
            else:
                stats["replaced"] += 1
    stats["edited"] = len(edited)
    # 4. Einfuegungen
    blocks.sort(key=lambda b: (b["order"], b["file"], b["idx"]))
    by_anchor = {}
    for b in blocks:
        for spec in b["slides"]:
            try:
                sl = build_slide(prs, tpl, spec, warns)
            except BuildError as exc:
                raise BuildError("%s: %s" % (spec["file"], exc))
            by_anchor.setdefault(b["after"], []).append(("new", sl))
            stats["inserted"] += 1
    final = []
    final.extend(by_anchor.get(0, []))
    for x in order:
        final.append(x)
        final.extend(by_anchor.get(x, []))
    new_id_by_slide = {}
    for el in list(sld_id_lst):
        for lst in by_anchor.values():
            for _, sl in lst:
                if int(el.get("id")) == sl.slide_id:
                    new_id_by_slide[id(sl)] = el
    # sldIdLst neu aufbauen
    for el in list(sld_id_lst):
        sld_id_lst.remove(el)
    for x in final:
        if isinstance(x, int):
            if x in deleted:
                continue
            sld_id_lst.append(orig_ids[x - 1])
        else:
            sld_id_lst.append(new_id_by_slide[id(x[1])])
    for num in deleted:
        prs.part.drop_rel(orig_ids[num - 1].get(qn("r:id")))
    stats["deleted"] = len(deleted)
    ext = prs.part._element.find(".//{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst")
    if ext is not None:
        warns.append("Praesentation enthaelt Abschnitte (sectionLst); neue/verschobene Folien sind evtl. nicht zugeordnet")
    total = len(prs.slides)  # benennt Folienparts gemaess Reihenfolge um
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)

    if not quiet:
        print("Eingabe : %s (%d Folien)" % (in_path, n))
        print("Inhalte : %d Datei(en)%s" % (len(files), "".join("\n          - " + os.path.basename(f) for f in files)))
        print("Ausgabe : %s" % out_path)
        print("Eingefuegt %d | verschoben %d | geloescht %d | editiert %d | Textersetzung auf %d Folie(n)"
              % (stats["inserted"], stats["moved"], stats["deleted"], stats["edited"], stats["replaced"]))
        print("Gesamtfolienzahl: %d" % total)
        for w in warns:
            print("WARNUNG: " + w)
    return stats, warns, total


def main(argv=None):
    ap = argparse.ArgumentParser(description="Neue Folien in presentation_angular21.pptx einfuegen")
    ap.add_argument("input", help="Original-PPTX (wird nicht veraendert)")
    ap.add_argument("output", help="Ausgabe-PPTX (wird neu geschrieben)")
    ap.add_argument("--content-dir", default=os.path.dirname(os.path.abspath(__file__)),
                    help="Ordner mit content_*.json (Standard: Ordner des Skripts)")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)
    try:
        build(args.input, args.output, args.content_dir, args.quiet)
    except BuildError as exc:
        print("FEHLER: %s" % exc, file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
