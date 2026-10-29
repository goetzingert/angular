"""Einfache Diagramme und Code-Kästen als native PowerPoint-Formen (für build_deck.py).

Eine Folie vom Typ `content` kann ein Feld `"diagram"` haben. Die Aufzählung steht dann links,
das Diagramm rechts. Unterstützte Arten (`kind`):

  flow     {"steps": ["A", "B", "C"], "note": "optional"}           Kästen untereinander mit Pfeilen
  layers   {"layers": ["oben", "mitte", "unten"], "note": ""}        gestapelte Schichten
  tree     {"root": "Eltern", "children": ["A", "B", "C"]}           ein Kasten, darunter Kinder
  compare  {"left": {"title": "", "items": []}, "right": {...}}     zwei Spalten
  cycle    {"steps": ["A", "B", "C"]}                                Kreislauf (3-4 Schritte, als Pfeilreihe im Quadrat)
  code     {"text": "zeilen\\nmit\\numbruch", "title": "Dateiname"}  Code-Kasten (Monospace)
  sequence {"actors": ["A", "B"], "steps": [{"from": 0, "to": 1, "text": "ruft auf"},
            {"from": 1, "to": 0, "text": "Ergebnis", "reply": true}]}       Sequenzdiagramm (Lebenslinien)
  composition {"root": {"label": "App", "children": [{"label": "Kind"}]},
            "links": [{"from": "A", "to": "B", "text": "[x]"}]}             verschachtelte Kästen (Komposition)

Mit Folientyp `diagram` steht das Diagramm über die ganze Folienbreite (Bereich WIDE).

Alle Maße in EMU. Der Bereich rechts ist X0..X0+W, Y0..Y0+H.
"""
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

X0, Y0, W, H = 7000000, 1850000, 4350000, 4350000
NARROW = (X0, Y0, W, H)
WIDE = (838080, 1700000, 10514160, 4650000)

GREEN = RGBColor(0x1E, 0x9E, 0x78)
DARK = RGBColor(0x2B, 0x3A, 0x42)
LIGHT = RGBColor(0xE3, 0xF4, 0xEE)
GREY = RGBColor(0xF2, 0xF2, 0xF2)
AMBER = RGBColor(0xFC, 0xE9, 0xD6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODEBG = RGBColor(0xF6, 0xF8, 0xFA)
FONT = "Segoe UI"


def _text(shape, text, size, color=DARK, bold=False, align=PP_ALIGN.CENTER, font=FONT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(70000)
    tf.margin_top = tf.margin_bottom = Emu(30000)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color


def _box(slide, x, y, w, h, text, fill=LIGHT, line=GREEN, size=14, color=DARK, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.18
    _text(s, text, size, color=color, bold=bold)
    return s


def _arrow(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y1), Emu(x2), Emu(y2))
    c.line.color.rgb = GREEN
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def _note(slide, text, y):
    s = slide.shapes.add_textbox(Emu(X0), Emu(y), Emu(W), Emu(380000))
    _text(s, text, 12, color=RGBColor(0x59, 0x59, 0x59), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    return s


def flow(slide, d):
    steps = d["steps"]
    n = len(steps)
    note_h = 450000 if d.get("note") else 0
    gap = 250000 if n <= 5 else 180000
    h = min(780000, int((H - note_h - gap * (n - 1)) / n))
    size = 17 if h > 600000 else 15
    for i, t in enumerate(steps):
        y = Y0 + i * (h + gap)
        last = i == n - 1
        _box(slide, X0, y, W, h, t, fill=GREEN if last and d.get("highlight_last") else LIGHT,
             color=WHITE if last and d.get("highlight_last") else DARK, size=size, bold=last and d.get("highlight_last", False))
        if not last:
            _arrow(slide, X0 + W // 2, y + h, X0 + W // 2, y + h + gap)
    if d.get("note"):
        _note(slide, d["note"], Y0 + n * h + (n - 1) * gap + 60000)


def layers(slide, d):
    items = d["layers"]
    n = len(items)
    note_h = 450000 if d.get("note") else 0
    gap = 90000
    h = min(900000, int((H - note_h - gap * (n - 1)) / n))
    fills = [GREEN, RGBColor(0x5B, 0xB9, 0x9C), RGBColor(0xA8, 0xDC, 0xCB), LIGHT, GREY]
    for i, t in enumerate(items):
        f = fills[min(i, len(fills) - 1)]
        col = WHITE if i < 2 else DARK
        _box(slide, X0, Y0 + i * (h + gap), W, h, t, fill=f, line=None, color=col, size=17, bold=i < 2)
    if d.get("note"):
        _note(slide, d["note"], Y0 + n * h + (n - 1) * gap + 60000)


def tree(slide, d):
    kids = d["children"]
    n = len(kids)
    cols = 2 if n > 3 else n
    rows = (n + cols - 1) // cols
    root_h = 700000
    gap_y = 450000
    gap_x = 150000
    cw = int((W - gap_x * (cols - 1)) / cols)
    ch = min(750000, int((H - root_h - gap_y - 120000 * (rows - 1)) / rows))
    _box(slide, X0 + W // 4, Y0, W // 2, root_h, d["root"], fill=GREEN, line=None, color=WHITE, size=16, bold=True)
    for i, t in enumerate(kids):
        r, c = divmod(i, cols)
        # letzte Zeile zentrieren
        in_row = cols if (r + 1) * cols <= n else n - r * cols
        total = in_row * cw + (in_row - 1) * gap_x
        x = X0 + (W - total) // 2 + c * (cw + gap_x)
        y = Y0 + root_h + gap_y + r * (ch + 120000)
        _box(slide, x, y, cw, ch, t, size=14)
        if r == 0:
            _arrow(slide, X0 + W // 2, Y0 + root_h, x + cw // 2, y)


def compare(slide, d):
    gap = 150000
    cw = (W - gap) // 2
    for i, key in enumerate(("left", "right")):
        side = d[key]
        x = X0 + i * (cw + gap)
        head = _box(slide, x, Y0, cw, 600000, side["title"], fill=GREEN if i == 0 else DARK, line=None, color=WHITE, size=15, bold=True)
        items = side["items"]
        ih = min(620000, int((H - 600000 - 100000 - 90000 * (len(items) - 1)) / max(1, len(items))))
        for j, t in enumerate(items):
            _box(slide, x, Y0 + 700000 + j * (ih + 90000), cw, ih, t, fill=LIGHT if i == 0 else GREY,
                 line=GREEN if i == 0 else RGBColor(0x8A, 0x96, 0x9C), size=14)


def cycle(slide, d):
    steps = d["steps"]
    n = len(steps)
    bw, bh = 1850000, 700000
    pos = [(X0, Y0 + 200000), (X0 + W - bw, Y0 + 200000), (X0 + W - bw, Y0 + 2400000), (X0, Y0 + 2400000)][:n]
    for t, (x, y) in zip(steps, pos):
        _box(slide, x, y, bw, bh, t, size=16)
    for i in range(n):
        (x1, y1), (x2, y2) = pos[i], pos[(i + 1) % n]
        if y1 == y2:
            if x2 > x1:
                _arrow(slide, x1 + bw, y1 + bh // 2, x2, y2 + bh // 2)
            else:
                _arrow(slide, x1, y1 + bh // 2, x2 + bw, y2 + bh // 2)
        else:
            if y2 > y1:
                _arrow(slide, x1 + bw // 2, y1 + bh, x2 + bw // 2, y2)
            else:
                _arrow(slide, x1 + bw // 2, y1, x2 + bw // 2, y2 + bh)
    if d.get("note"):
        _note(slide, d["note"], Y0 + 1550000)


def code(slide, d):
    lines = d["text"].split("\n")
    size = 15 if len(lines) <= 9 and max(len(l) for l in lines) <= 34 else 13
    title = d.get("title")
    top = Y0
    if title:
        _box(slide, X0, top, W, 380000, title, fill=DARK, line=None, color=WHITE, size=12, shape=MSO_SHAPE.RECTANGLE)
        top += 380000
    h = min(H - (380000 if title else 0), int(len(lines) * size * 1.35 * 12700) + 220000)
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(X0), Emu(top), Emu(W), Emu(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = CODEBG
    s.line.color.rgb = RGBColor(0xC8, 0xD0, 0xD6)
    s.line.width = Pt(1)
    _text(s, d["text"], size, color=DARK, align=PP_ALIGN.LEFT, font="Consolas", anchor=MSO_ANCHOR.TOP)
    s.text_frame.margin_left = Emu(120000)
    s.text_frame.margin_top = Emu(90000)
    if d.get("note"):
        _note(slide, d["note"], top + h + 60000)


def _line(slide, x1, y1, x2, y2, color=GREEN, width=2, dashed=False, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y1), Emu(x2), Emu(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def _label(slide, x, y, w, h, text, size, color=DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, bold=False,
           bg=False):
    s = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    _text(s, text, size, color=color, align=align, anchor=anchor, bold=bold)
    if bg:
        s.fill.solid()
        s.fill.fore_color.rgb = WHITE
    s.text_frame.margin_top = s.text_frame.margin_bottom = Emu(0)
    return s


def sequence(slide, d):
    """Sequenzdiagramm: Teilnehmer oben, gestrichelte Lebenslinien, nummerierte Nachrichten."""
    actors, steps = d["actors"], d["steps"]
    n, m = len(actors), len(steps)
    note_h = 420000 if d.get("note") else 0
    cw = W // n
    hw, hh = min(cw - 90000, 2300000), 560000
    wide = W > 6000000
    size = (15 if n <= 4 else 13) if wide else 12
    hsize = (16 if n <= 4 else 14) if wide else 13
    xs = [X0 + cw * i + cw // 2 for i in range(n)]
    for i, a in enumerate(actors):
        _box(slide, xs[i] - hw // 2, Y0, hw, hh, a, fill=DARK, line=None, color=WHITE, size=hsize, bold=True,
             shape=MSO_SHAPE.RECTANGLE)
    top = Y0 + hh
    bottom = Y0 + H - note_h
    pitch = min(620000, (bottom - top - 120000) // m)
    end = min(bottom, top + 120000 + pitch * m + 120000)
    for x in xs:
        _line(slide, x, top, x, end, color=RGBColor(0x9A, 0xA5, 0xAB), width=1.25, dashed=True, arrow=False)
    lab_h = min(330000, pitch - 60000)
    for k, st in enumerate(steps):
        y = top + 120000 + pitch * k + int(pitch * 0.72)
        a, b = st["from"], st["to"]
        text = "%d. %s" % (k + 1, st["text"])
        reply = st.get("reply", False)
        col = RGBColor(0x5B, 0x6B, 0x73) if reply else GREEN
        if a == b:
            loop_w, loop_h = min(380000, cw // 4), int(pitch * 0.38)
            x = xs[a]
            lw = max(cw - loop_w - 120000, 900000)
            left_side = a == n - 1 and n > 1   # letzter Teilnehmer: Schleife und Text nach links
            sgn = -1 if left_side else 1
            _line(slide, x, y - loop_h, x + sgn * loop_w, y - loop_h, color=col, width=1.75, arrow=False)
            _line(slide, x + sgn * loop_w, y - loop_h, x + sgn * loop_w, y, color=col, width=1.75, arrow=False)
            _line(slide, x + sgn * loop_w, y, x + sgn * 20000, y, color=col, width=1.75)
            if left_side:
                lx = x - loop_w - 50000 - lw
                al = PP_ALIGN.RIGHT
            else:
                lx = x + loop_w + 50000
                al = PP_ALIGN.LEFT
            _label(slide, lx, y - loop_h - lab_h // 2, lw, lab_h, text, size, align=al, anchor=MSO_ANCHOR.MIDDLE)
        else:
            x1, x2 = xs[a], xs[b]
            _line(slide, x1, y, x2 + (-20000 if x2 > x1 else 20000), y, color=col, width=1.75, dashed=reply)
            width = max(abs(x2 - x1) - 60000, min(int(cw * 1.5), W - 40000))
            left = max(X0, min((x1 + x2) // 2 - width // 2, X0 + W - width))
            _label(slide, left, y - lab_h - 20000, width, lab_h, text, size,
                   color=RGBColor(0x40, 0x4B, 0x52) if reply else DARK)
    if d.get("note"):
        _note(slide, d["note"], end + 60000)


def _comp_count(node):
    return 1 + sum(_comp_count(k) for k in node.get("children", []))


COMP_FILLS = [RGBColor(0xFA, 0xFB, 0xFC), LIGHT, WHITE, RGBColor(0xF2, 0xF7, 0xF5)]
LEAF_H, LEAF_SUB_H, PAD, GAP_V = 640000, 820000, 110000, 110000


def _head(node):
    return 330000 + (230000 if node.get("sub") else 0)


def _nat_h(node, gap_v):
    kids = node.get("children", [])
    if not kids:
        return LEAF_SUB_H if node.get("sub") else LEAF_H
    hs = [_nat_h(k, gap_v) for k in kids]
    inner = sum(hs) + gap_v * (len(kids) - 1) if node.get("vertical") else max(hs)
    return _head(node) + inner + PAD


def _comp_node(slide, node, x, y, w, h, depth, boxes, size, scale, gap_h, gap_v):
    kids = node.get("children", [])
    hl = node.get("highlight", False)
    label = node["label"] + ("\n" + node["sub"] if node.get("sub") else "")
    if kids:
        s = _box(slide, x, y, w, h, "", fill=COMP_FILLS[min(depth, 3)], line=GREEN if depth < 2 else RGBColor(0x8A, 0x96, 0x9C),
                 size=size, shape=MSO_SHAPE.RECTANGLE)
        s.line.width = Pt(1.5 if depth == 0 else 1)
        head = int(_head(node) * scale)
        _label(slide, x + 60000, y + 40000, w - 120000, head - 40000, label, size, bold=True, align=PP_ALIGN.LEFT,
               anchor=MSO_ANCHOR.TOP)
        pad = PAD
        inner_y, inner_h = y + head, h - head - int(pad * scale)
        pad_v = int(pad * scale)  # noqa: F841 (Abstand unten skaliert, damit nichts überläuft)
        n = len(kids)
        if node.get("vertical"):
            g = int(gap_v * scale)
            hs = [int(_nat_h(k, gap_v) * scale) for k in kids]
            extra = max(0, inner_h - sum(hs) - g * (n - 1))
            yy = inner_y
            for k, kh in zip(kids, hs):
                kh2 = kh + (extra // n if k.get("children") else 0)
                _comp_node(slide, k, x + pad, yy, w - 2 * pad, kh2, depth + 1, boxes, size, scale, gap_h, gap_v)
                yy += kh2 + g
        else:
            cw = (w - 2 * pad - gap_h * (n - 1)) // n
            for i, k in enumerate(kids):
                kh = inner_h if k.get("children") else int(_nat_h(k, gap_v) * scale)
                _comp_node(slide, k, x + pad + i * (cw + gap_h), inner_y, cw, kh, depth + 1, boxes, size, scale, gap_h, gap_v)
    else:
        _box(slide, x, y, w, h, label, fill=GREEN if hl else LIGHT, color=WHITE if hl else DARK, size=size, bold=hl)
    boxes[node["label"]] = (x, y, w, h)


def composition(slide, d):
    """Verschachtelte Kästen: wer enthält wen (Template-Komposition, Provider-Baum ...)."""
    note_h = 420000 if d.get("note") else 0
    links = d.get("links", [])
    size = d.get("size", 13 if W > 6000000 else 12)
    boxes = {}
    gap_h = 520000 if links else 140000
    gap_v = 420000 if links else GAP_V
    avail = H - note_h
    nat = _nat_h(d["root"], gap_v)
    scale = min(1.0, avail / nat)
    root_h = int(nat * scale) if d["root"].get("children") else min(avail, LEAF_H)
    _comp_node(slide, d["root"], X0, Y0, W, root_h, 0, boxes, size, scale, gap_h, gap_v)
    for ln in links:
        if ln["from"] not in boxes or ln["to"] not in boxes:
            continue
        ax, ay, aw, ah = boxes[ln["from"]]
        bx, by, bw, bh = boxes[ln["to"]]
        acx, acy, bcx, bcy = ax + aw // 2, ay + ah // 2, bx + bw // 2, by + bh // 2
        if abs(bcx - acx) >= abs(bcy - acy):
            x1 = ax + aw if bcx > acx else ax
            x2 = bx if bcx > acx else bx + bw
            y1 = y2 = (max(ay, by) + min(ay + ah, by + bh)) // 2 if not (ay + ah < by or by + bh < ay) else acy
            if y1 != acy:
                y2 = y1
            else:
                y2 = bcy
        else:
            y1 = ay + ah if bcy > acy else ay
            y2 = by if bcy > acy else by + bh
            x1 = x2 = (max(ax, bx) + min(ax + aw, bx + bw)) // 2 if not (ax + aw < bx or bx + bw < ax) else acx
            if x1 == acx:
                x2 = bcx
        _line(slide, x1, y1, x2, y2, color=RGBColor(0xC0, 0x5A, 0x1E) if ln.get("accent") else GREEN, width=2,
              dashed=ln.get("dashed", False))
        if ln.get("text"):
            mx, my = (x1 + x2) // 2, (y1 + y2) // 2
            col = RGBColor(0xC0, 0x5A, 0x1E) if ln.get("accent") else GREEN
            # Beschriftung mit weißem Hintergrund, damit sie auch über Kastenrändern lesbar bleibt
            chars = max(len(t) for t in ln["text"].split("\n"))
            lw = int(chars * (size - 1) * 0.62 * 12700) + 160000
            if x1 == x2:   # senkrecht: Beschriftung rechts daneben
                _label(slide, mx + 60000, my - 140000, lw, 280000, ln["text"], size - 1, color=col, bold=True,
                       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, bg=True)
            else:          # waagerecht: Beschriftung mittig über dem Pfeil
                _label(slide, mx - lw // 2, my - 330000, lw, 280000, ln["text"], size - 1, color=col, bold=True,
                       anchor=MSO_ANCHOR.MIDDLE, bg=True)
    if d.get("note"):
        _note(slide, d["note"], Y0 + root_h + 60000)


KINDS = {"flow": flow, "layers": layers, "tree": tree, "compare": compare, "cycle": cycle, "code": code,
         "sequence": sequence, "composition": composition}


def validate(d, where, wide=False):
    """Gibt eine Fehlermeldung oder None zurück."""
    if not isinstance(d, dict) or d.get("kind") not in KINDS:
        return "%s: diagram.kind muss eine von %s sein" % (where, ", ".join(KINDS))
    k = d["kind"]
    need = {"flow": "steps", "layers": "layers", "tree": "children", "compare": "left", "cycle": "steps", "code": "text",
            "sequence": "steps", "composition": "root"}[k]
    if need not in d:
        return "%s: diagram (%s) braucht '%s'" % (where, k, need)
    if k == "tree" and "root" not in d:
        return "%s: diagram (tree) braucht 'root'" % where
    if k == "compare" and "right" not in d:
        return "%s: diagram (compare) braucht 'right'" % where
    if k == "cycle" and not 3 <= len(d["steps"]) <= 4:
        return "%s: diagram (cycle) braucht 3 oder 4 Schritte" % where
    if k == "flow" and not 2 <= len(d["steps"]) <= 6:
        return "%s: diagram (flow) braucht 2 bis 6 Schritte" % where
    if k == "layers" and not 2 <= len(d["layers"]) <= 6:
        return "%s: diagram (layers) braucht 2 bis 6 Ebenen" % where
    if k == "tree" and not 2 <= len(d["children"]) <= 6:
        return "%s: diagram (tree) braucht 2 bis 6 Kinder" % where
    if k == "code" and len(d["text"].split("\n")) > 16:
        return "%s: diagram (code) höchstens 16 Zeilen" % where
    if k == "sequence":
        a = d.get("actors")
        maxa = 6 if wide else 3
        if not isinstance(a, list) or not 2 <= len(a) <= maxa:
            return "%s: diagram (sequence) braucht 2 bis %d actors (%s)" % (where, maxa, "Folientyp diagram" if wide else "neben Text; mehr mit type diagram")
        st = d["steps"]
        maxs = 10 if wide else 8
        if not isinstance(st, list) or not 1 <= len(st) <= maxs:
            return "%s: diagram (sequence) braucht 1 bis %d steps" % (where, maxs)
        for i, x in enumerate(st):
            if not isinstance(x, dict) or not isinstance(x.get("text"), str) or \
                    x.get("from") not in range(len(a)) or x.get("to") not in range(len(a)):
                return "%s: diagram (sequence) steps[%d] braucht from/to (Index in actors) und text" % (where, i)
    if k == "composition":
        def chk(node, depth):
            if not isinstance(node, dict) or not isinstance(node.get("label"), str):
                return "Knoten ohne label"
            kids = node.get("children", [])
            if depth > 3:
                return "höchstens 4 Ebenen"
            if len(kids) > (5 if node.get("vertical") else 4):
                return "höchstens 4 Kinder nebeneinander (5 mit vertical)"
            for c in kids:
                e = chk(c, depth + 1)
                if e:
                    return e
            return None
        e = chk(d["root"], 0)
        if e:
            return "%s: diagram (composition): %s" % (where, e)
        if _comp_count(d["root"]) > 16:
            return "%s: diagram (composition): höchstens 16 Kästen" % where
        labels = []
        def collect(node):
            labels.append(node["label"])
            for c in node.get("children", []):
                collect(c)
        collect(d["root"])
        if len(set(labels)) != len(labels):
            return "%s: diagram (composition): labels müssen eindeutig sein (für links)" % where
        for i, ln in enumerate(d.get("links", [])):
            if ln.get("from") not in labels or ln.get("to") not in labels:
                return "%s: diagram (composition) links[%d]: from/to müssen labels sein" % (where, i)
    return None


def caption(slide, text, box):
    """Einzeilige Bildunterschrift unter einem breiten Diagramm."""
    x, y, w, h = box
    _label(slide, x, y, w, h, text, 16, color=RGBColor(0x40, 0x4B, 0x52), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)


def draw(slide, d, region=None):
    global X0, Y0, W, H
    saved = (X0, Y0, W, H)
    X0, Y0, W, H = region or NARROW
    try:
        KINDS[d["kind"]](slide, d)
    finally:
        X0, Y0, W, H = saved
